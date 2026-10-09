import unittest
from io import BytesIO

import pymupdf
from docx import Document

from app import app, format_extracted_preview
from nlp.summarizer import summarize_text


SAMPLE_TEXT = (
    "The coastal research team measured water quality across the bay. "
    "The results showed that runoff increased nutrient levels after storms. "
    "Higher nutrient levels caused rapid algae growth in several locations. "
    "The team recommends reducing fertilizer use near waterways. "
    "Local wetlands can help filter pollutants before they reach the bay. "
    "Follow-up testing will measure whether these changes improve water quality."
)


class DocumentAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_summary_returns_selected_sentences_in_document_order(self):
        summary = summarize_text(SAMPLE_TEXT, sentence_count=4)

        self.assertEqual(len(summary), 4)
        positions = [SAMPLE_TEXT.index(sentence) for sentence in summary]
        self.assertEqual(positions, sorted(positions))

    def test_extracted_preview_joins_standalone_bullet_to_text(self):
        extracted_text = (
            "6. Open Problems\n"
            "•\n\n"
            "Security: Replay weaknesses are documented.\n"
            "•\n"
            "Privacy: Payment metadata can leak information."
        )

        self.assertEqual(
            format_extracted_preview(extracted_text),
            "6. Open Problems\n"
            "• Security: Replay weaknesses are documented.\n"
            "• Privacy: Payment metadata can leak information."
        )

    def test_docx_upload_extracts_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph(SAMPLE_TEXT)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Monitoring site"
        table.cell(0, 1).text = "North Bay"
        document_buffer = BytesIO()
        document.save(document_buffer)

        response = self.client.post(
            "/",
            data={
                "document": (
                    BytesIO(document_buffer.getvalue()),
                    "research.docx"
                )
            },
            content_type="multipart/form-data"
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Document summary", response.data)
        self.assertIn(b"North Bay", response.data)
        self.assertRegex(
            response.data,
            rb"<span>Pages</span>\s*<strong>\s*\xe2\x80\x94\s*</strong>"
        )
        self.assertEqual(response.data.count(b"<li>"), 5)

    def test_pdf_upload_keeps_page_count_and_summary(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), SAMPLE_TEXT)
        pdf_bytes = document.tobytes()
        document.close()

        response = self.client.post(
            "/",
            data={"document": (BytesIO(pdf_bytes), "research.pdf")},
            content_type="multipart/form-data"
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Document summary", response.data)
        self.assertRegex(
            response.data,
            rb"<span>Pages</span>\s*<strong>\s*1\s*</strong>"
        )

    def test_unsupported_file_type_shows_validation_error(self):
        response = self.client.post(
            "/",
            data={"document": (BytesIO(b"not a document"), "notes.txt")},
            content_type="multipart/form-data"
        )

        self.assertIn(b"Only PDF and Word", response.data)


if __name__ == "__main__":
    unittest.main()
